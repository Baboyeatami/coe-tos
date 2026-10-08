"""Local extraction with stable source locators; no model/API dependency."""
from pathlib import Path
import shutil
import subprocess
import tempfile

import openpyxl
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from office import input_backend, office_binary


def convert_office(source, extension, destination):
    binary = office_binary()
    if not binary:
        raise ValueError("Legacy Office conversion needs LibreOffice; convert to DOCX/XLSX first.")
    with tempfile.TemporaryDirectory(prefix="coe-tos-lo-") as profile:
        subprocess.run([binary, f"-env:UserInstallation={Path(profile).as_uri()}",
                        "--headless", "--convert-to", extension,
                        "--outdir", str(destination), str(source)],
                       check=True, capture_output=True, text=True, timeout=180)
    result = Path(destination) / (source.stem + "." + extension.split(":")[0])
    if not result.is_file():
        raise ValueError(f"LibreOffice did not create {result.name}")
    return result


def text_document(text, name, warnings=None):
    """Document for text supplied directly in chat, using the same line locators as a file."""
    result = {"source": name, "format": ".txt", "segments": [], "warnings": []}
    for line_number, line in enumerate(text.splitlines(), 1):
        if str(line).strip():
            result["segments"].append({"locator": f"line {line_number}", "text": str(line)})
    if not result["segments"]:
        result["warnings"].append("No readable content supplied; do not draft from an empty source.")
    else:
        result["warnings"].append("Text was supplied directly in chat: table layout, emphasis and totals cannot be visually verified. Read published scores back to the requester for confirmation.")
    result["warnings"].extend(warnings or [])
    return result


def extract(path, ocr=False, office_backend="auto"):
    path = Path(path).resolve()
    if not path.is_file():
        raise ValueError(f"Input is not a file: {path}")
    result = {"source": path.name, "format": path.suffix.lower(),
              "segments": [], "warnings": []}
    segments = result["segments"]

    def add(locator, text, **extra):
        if str(text).strip():
            segments.append({"locator": locator, "text": str(text), **extra})

    suffix = path.suffix.lower()
    if suffix in (".doc", ".xls", ".odt", ".ods", ".rtf"):
        extension = "xlsx" if suffix in (".xls", ".ods") else "docx"
        backend = input_backend(suffix, office_backend)
        with tempfile.TemporaryDirectory(prefix="coe-tos-input-") as directory:
            if backend == "libreoffice":
                converted = convert_office(path, extension, directory)
                converter = "LibreOffice"
            else:
                from windows_office import run_office
                # Converters see copies of legacy inputs as well as output workbooks.
                copied = Path(directory) / ("input" + suffix)
                shutil.copy2(path, copied)
                converted = run_office(extension, copied, Path(directory) / ("converted." + extension))
                converter = "Microsoft " + ("Word" if backend == "word-windows" else "Excel")
            child = extract(converted, ocr=ocr, office_backend=office_backend)
            result.update(segments=child["segments"], warnings=child["warnings"])
        result["warnings"].append(f"Converted through {converter}; locators refer to converted content.")
    elif suffix == ".pdf":
        if ocr:
            binary = shutil.which("ocrmypdf")
            if not binary:
                raise ValueError("--ocr on PDFs requires OCRmyPDF (and its system dependencies).")
            with tempfile.TemporaryDirectory(prefix="coe-tos-ocr-") as directory:
                target = Path(directory) / "ocr.pdf"
                subprocess.run([binary, "--skip-text", str(path), str(target)],
                               check=True, capture_output=True, timeout=600)
                child = extract(target)
                result.update(segments=child["segments"], warnings=child["warnings"])
            result["warnings"].append("OCR used; verify formulas, symbols and scores against page images.")
        else:
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("PDF is encrypted; supply an unlocked copy.")
            for page_number, page in enumerate(reader.pages, 1):
                text = (page.extract_text(extraction_mode="layout") or "") if "/Contents" in page else ""
                add(f"page {page_number}", text)
                if len(text.strip()) < 30:
                    result["warnings"].append(f"Page {page_number}: little/no extractable text; inspect visually or use --ocr.")
            result["warnings"].append("PDF tables/layout require visual review; extraction is page text, not guaranteed table reconstruction.")
    elif suffix == ".docx":
        doc = Document(path)
        for index, element in enumerate(doc.element.body, 1):
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(element, doc)
                add(f"body {index}", paragraph.text, style=paragraph.style.name)
            elif tag == "tbl":
                table = Table(element, doc)
                for row_index, row in enumerate(table.rows, 1):
                    for col_index, cell in enumerate(row.cells, 1):
                        add(f"table body {index}, row {row_index}, column {col_index}", cell.text)
        for index, section in enumerate(doc.sections, 1):
            for name in ("header", "footer"):
                for paragraph in getattr(section, name).paragraphs:
                    add(f"section {index} {name}", paragraph.text)
        result["warnings"].append("Floating text boxes, equations and image-only content need visual review.")
    elif suffix in (".xlsx", ".xlsm"):
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in formulas:
                for row in sheet:
                    for cell in row:
                        if cell.value is None:
                            continue
                        extra = {}
                        if cell.data_type == "f":
                            value = cached[sheet.title][cell.coordinate].value
                            extra = {"formula": cell.value, "cached_value": value}
                            if value is None:
                                result["warnings"].append(f"{sheet.title}!{cell.coordinate}: formula has no stored result.")
                        add(f"{sheet.title}!{cell.coordinate}", cell.value, **extra)
        finally:
            formulas.close()
            cached.close()
    elif suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
        binary = shutil.which("tesseract")
        if not ocr or not binary:
            raise ValueError("Images need --ocr with Tesseract, or the agent's visual analysis tool.")
        process = subprocess.run([binary, str(path), "stdout"], capture_output=True,
                                 text=True, check=True, timeout=180)
        add("image OCR", process.stdout)
        result["warnings"].append("OCR used; inspect symbols and tables against the image.")
    else:
        # Text formats and unknown files that actually contain UTF text.
        raw = path.read_bytes()
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        try:
            text = raw.decode(encoding)
        except UnicodeError as error:
            raise ValueError(f"Unsupported binary format {suffix or '(no extension)'}; use a harness extractor or convert to a supported format.") from error
        if "\x00" in text:
            raise ValueError("Input appears binary; use an appropriate extractor or convert it first.")
        for line_number, line in enumerate(text.splitlines(), 1):
            add(f"line {line_number}", line)
    if not segments:
        result["warnings"].append("No readable content extracted; do not draft from an empty source.")
    return result
