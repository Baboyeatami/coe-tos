"""Edit only worksheet values/caches and explicit print settings in an XLSX ZIP."""
from io import BytesIO
from pathlib import Path
import hashlib
import posixpath
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl
from openpyxl.utils import coordinate_to_tuple

from formulas import evaluator
from scoring import GROUPS, allocation_index, number

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
Q = lambda name: f"{{{NS}}}{name}"
QCT = lambda name: f"{{{CT}}}{name}"


def parse(data):
    return ET.fromstring(data, ET.XMLParser(resolve_entities=False, no_network=True))


def serialize(root):
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def package(path):
    with ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("Workbook ZIP integrity check failed")
        return {n: archive.read(n) for n in archive.namelist()}


def sheet_paths(parts):
    workbook = parse(parts["xl/workbook.xml"])
    rels = {r.get("Id"): r.get("Target")
            for r in parse(parts["xl/_rels/workbook.xml.rels"])}
    result = {}
    for sheet in workbook.find(Q("sheets")):
        target = rels[sheet.get(f"{{{REL}}}id")]
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        result[sheet.get("name")] = path
    return result


def inspect(path):
    parts = package(path)
    wb = openpyxl.load_workbook(BytesIO(Path(path).read_bytes()), data_only=False)
    result = {"sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
              "sheets": [], "drawing_and_media_parts": sorted(
                  n for n in parts if n.startswith(("xl/drawings/", "xl/media/")))}
    for sheet in wb:
        cells = []
        for row in sheet:
            for cell in row:
                if cell.value is not None or cell.comment:
                    cells.append({"cell": cell.coordinate, "value": cell.value,
                                  "formula": cell.data_type == "f",
                                  "comment": cell.comment.text if cell.comment else None,
                                  "style_id": cell.style_id, "number_format": cell.number_format})
        result["sheets"].append({"name": sheet.title, "cells": cells,
                                 "merges": list(map(str, sheet.merged_cells.ranges)),
                                 "row_heights": {str(k): v.height for k, v in sheet.row_dimensions.items()},
                                 "columns": {k: {"min": v.min, "max": v.max, "width": v.width}
                                             for k, v in sheet.column_dimensions.items()},
                                 "validations": [{"type": v.type, "cells": str(v.sqref),
                                                  "formula1": v.formula1, "formula2": v.formula2}
                                                 for v in sheet.data_validations.dataValidation],
                                 "print_area": str(sheet.print_area),
                                 "orientation": sheet.page_setup.orientation,
                                 "paper_size": sheet.page_setup.paperSize})
    return result


def cell_node(root, address):
    data = root.find(Q("sheetData"))
    row_index, column_index = coordinate_to_tuple(address)
    row = next((r for r in data if int(r.get("r")) == row_index), None)
    if row is None:
        row = ET.Element(Q("row"), r=str(row_index))
        before = next((i for i, r in enumerate(data) if int(r.get("r")) > row_index), len(data))
        data.insert(before, row)
    node = next((c for c in row if c.get("r") == address), None)
    if node is None:
        node = ET.Element(Q("c"), r=address)
        before = next((i for i, c in enumerate(row)
                       if c.tag == Q("c") and coordinate_to_tuple(c.get("r"))[1] > column_index), len(row))
        row.insert(before, node)
    return node


def write_value(node, value):
    if node.find(Q("f")) is not None:
        raise ValueError(f"Refusing to overwrite template formula at {node.get('r')}")
    for name in ("v", "is"):
        for child in node.findall(Q(name)):
            node.remove(child)
    node.attrib.pop("t", None)
    if value is None:
        return
    if isinstance(value, str):
        node.set("t", "inlineStr")
        text = ET.SubElement(ET.SubElement(node, Q("is")), Q("t"))
        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text.text = value
    else:
        ET.SubElement(node, Q("v")).text = str(value)


def write_cache(node, value):
    for child in node.findall(Q("v")):
        node.remove(child)
    node.attrib.pop("t", None)
    if value is None:
        return
    if isinstance(value, bool):
        node.set("t", "b")
        text = str(int(value))
    elif isinstance(value, str):
        node.set("t", "str")
        text = value
    else:
        text = str(value)
    ET.SubElement(node, Q("v")).text = text


def field_values(draft, profile):
    values = {}
    for name, address in profile.get("fields", {}).items():
        value = draft.get("metadata", {}).get(name)
        # Missing information stays visibly unresolved rather than inheriting an old year/name.
        values[address] = profile.get("prefixes", {}).get(name, "") + (str(value) if value else "[To be supplied]")
    return values


def fit_refs(refs, limit):
    """Abbreviate a Test Item No. cell that cannot hold every reference.

    The institutional row height is fixed, so a long list wraps and is clipped.
    Keep as many references as fit and state how many were dropped; the full list
    stays in the allocation ledger and the supplementary mapping sheets.
    """
    if not refs:
        return ""
    joined = "; ".join(refs)
    if not limit or len(joined) <= limit:
        return joined
    kept = 1
    while kept < len(refs) - 1:
        kept += 1
        candidate = "; ".join(refs[:kept]) + f" +{len(refs) - kept}"
        if len(candidate) > limit:
            kept -= 1
            break
    return "; ".join(refs[:kept]) + f" +{len(refs) - kept}"


def projected_values(draft, profile, result, index=None):
    lookup = index or allocation_index(draft)
    values = field_values(draft, profile)
    columns = profile["columns"]
    for index in range(profile["last_row"] - profile["first_row"] + 1):
        row = profile["first_row"] + index
        topic = draft["topics"][index] if index < len(draft["topics"]) else None
        scores = result["topic_scores"][topic["id"]] if topic else None
        basic = {"number": index + 1 if topic else None, "topic": topic["title"] if topic else None}
        for group in GROUPS:
            refs = sorted(lookup["references"].get((topic["id"], group), ())) if topic else []
            basic[group + "_refs"] = fit_refs(refs, profile.get("refs_max_chars")) if refs else ("--" if topic else None)
            basic[group] = scores[group] if scores else None
        for name, value in basic.items():
            values[f"{columns[name]}{row}"] = value
    return values


def computed_values(draft, profile, result):
    """Fill numeric summaries when a different template has no formulas there."""
    values = {}
    columns = profile["columns"]
    total = result["total"]
    for index in range(profile["last_row"] - profile["first_row"] + 1):
        row = profile["first_row"] + index
        topic = draft["topics"][index] if index < len(draft["topics"]) else None
        scores = result["topic_scores"][topic["id"]] if topic else None
        points = sum(scores.values()) if scores else None
        for group in GROUPS:
            values[f"{columns[group + '_percent']}{row}"] = scores[group] / total if scores else None
        values[f"{columns['points']}{row}"] = points
        values[f"{columns['percent']}{row}"] = points / total if points is not None else None
    row = profile["total_row"]
    for group in GROUPS:
        values[f"{columns[group]}{row}"] = result["cognitive"][group]
        values[f"{columns[group + '_percent']}{row}"] = result["cognitive"][group] / total
    values[f"{columns['points']}{row}"] = total
    values[f"{columns['percent']}{row}"] = 1
    return values


def apply_print(parts, paths, profile):
    sheet = parse(parts[paths[profile["sheet"]]])
    if profile.get("fit_to_page"):
        prop = sheet.find(Q("sheetPr"))
        if prop is None:
            prop = ET.Element(Q("sheetPr"))
            sheet.insert(0, prop)
        page = prop.find(Q("pageSetUpPr"))
        if page is None:
            page = ET.SubElement(prop, Q("pageSetUpPr"))
        page.set("fitToPage", "1")
        setup = sheet.find(Q("pageSetup"))
        if setup is None:
            raise ValueError("Template lacks pageSetup; define print settings in Excel first")
        setup.set("fitToWidth", "1")
        setup.set("fitToHeight", "1")
    parts[paths[profile["sheet"]]] = serialize(sheet)
    workbook = parse(parts["xl/workbook.xml"])
    index = list(paths).index(profile["sheet"])
    if profile.get("print_area"):
        names = workbook.find(Q("definedNames"))
        if names is None:
            names = ET.Element(Q("definedNames"))
            sheets = workbook.find(Q("sheets"))
            workbook.insert(list(workbook).index(sheets) + 1, names)
        name = next((n for n in names if n.get("name") == "_xlnm.Print_Area"
                     and n.get("localSheetId") == str(index)), None)
        if name is None:
            name = ET.SubElement(names, Q("definedName"), name="_xlnm.Print_Area", localSheetId=str(index))
        sheet_name = profile["sheet"].replace("'", "''")
        name.text = f"'{sheet_name}'!{profile['print_area']}"
    calc = workbook.find(Q("calcPr"))
    if calc is None:
        calc = ET.SubElement(workbook, Q("calcPr"))
    calc.set("fullCalcOnLoad", "1")
    calc.set("forceFullCalc", "1")
    calc.set("calcMode", "auto")
    parts["xl/workbook.xml"] = serialize(workbook)


def normalized_sheet(data, edited, print_changes=False):
    root = parse(data)
    for node in root.findall(f".//{Q('sheetData')}/{Q('row')}/{Q('c')}"):
        address = node.get("r")
        if address in edited:
            node.getparent().remove(node)
        elif node.find(Q("f")) is not None:
            node.attrib.pop("t", None)
            for child in node.findall(Q("v")):
                node.remove(child)
    for row in root.find(Q("sheetData")):
        if not len(row) and set(row.attrib) == {"r"}:
            row.getparent().remove(row)
    if print_changes:
        setup = root.find(Q("pageSetup"))
        if setup is not None:
            for key in ("fitToWidth", "fitToHeight"):
                setup.attrib.pop(key, None)
        prop = root.find(Q("sheetPr"))
        if prop is not None:
            page = prop.find(Q("pageSetUpPr"))
            if page is not None:
                page.attrib.pop("fitToPage", None)
                if not page.attrib and not len(page):
                    prop.remove(page)
            if not prop.attrib and not len(prop):
                root.remove(prop)
    return ET.tostring(root, method="c14n")


def verify(template, output, profile, edited, added=None, original_parts=None):
    added = added or {"parts": [], "relationship_ids": [], "sheet_names": []}
    new_parts = set(added["parts"])
    relationship_ids = set(added["relationship_ids"])
    sheet_names = set(added["sheet_names"])
    before = original_parts if original_parts is not None else package(template)
    after = package(output)
    if set(after) - set(before) != new_parts or set(before) - set(after):
        raise ValueError("Workbook package parts added or dropped")
    paths = sheet_paths(before)
    for name in before:
        if name == "xl/workbook.xml":
            a, b = parse(before[name]), parse(after[name])
            for root in (a, b):
                for calc in root.findall(Q("calcPr")):
                    root.remove(calc)
                names = root.find(Q("definedNames"))
                if names is not None:
                    for n in list(names):
                        if n.get("name") == "_xlnm.Print_Area" and n.get("localSheetId") == str(list(paths).index(profile["sheet"])):
                            names.remove(n)
                    if not len(names):
                        root.remove(names)
            # Declared supplementary sheets are additions, not edits to existing sheets.
            for node in list(b.find(Q("sheets"))):
                if node.get("name") in sheet_names:
                    b.find(Q("sheets")).remove(node)
            if ET.tostring(a, method="c14n") != ET.tostring(b, method="c14n"):
                raise ValueError("Workbook properties beyond print area/calculation changed")
        elif name == "[Content_Types].xml":
            a, b = parse(before[name]), parse(after[name])
            for node in list(b.findall(QCT("Override"))):
                if node.get("PartName", "").lstrip("/") in new_parts:
                    b.remove(node)
            if ET.tostring(a, method="c14n") != ET.tostring(b, method="c14n"):
                raise ValueError("Content types changed beyond declared supplementary sheets")
        elif name == "xl/_rels/workbook.xml.rels":
            a, b = parse(before[name]), parse(after[name])
            for node in list(b):
                if node.get("Id") in relationship_ids:
                    b.remove(node)
            if ET.tostring(a, method="c14n") != ET.tostring(b, method="c14n"):
                raise ValueError("Workbook relationships changed beyond declared supplementary sheets")
        elif name in paths.values():
            target = name == paths[profile["sheet"]]
            if normalized_sheet(before[name], edited if target else set(), target) != normalized_sheet(after[name], edited if target else set(), target):
                raise ValueError(f"Worksheet formatting/formulas changed: {name}")
            if target:
                a, b = parse(before[name]), parse(after[name])
                old = {c.get("r"): c for c in a.findall(f".//{Q('c')}")}
                new = {c.get("r"): c for c in b.findall(f".//{Q('c')}")}
                for address in edited:
                    old_attrs = {k: v for k, v in old[address].attrib.items() if k != "t"} if address in old else {"r": address}
                    new_attrs = {k: v for k, v in new[address].attrib.items() if k != "t"}
                    if old_attrs != new_attrs:
                        raise ValueError(f"Cell formatting changed: {address}")
        elif before[name] != after[name]:
            raise ValueError(f"Template part changed: {name}")
    for part in new_parts:
        if not after.get(part):
            raise ValueError(f"Declared supplementary sheet is missing or empty: {part}")
    return {"package_parts_preserved": len(before), "formula_and_style_fidelity": True,
            "unchanged_drawings_and_media": len([n for n in before if n.startswith(("xl/drawings/", "xl/media/"))]),
            "added_sheets": sorted(sheet_names), "added_parts": sorted(new_parts)}


def build(template, output, draft, profile, result, extra_sheets=None, hidden_sheets=True, index=None):
    template, output = Path(template), Path(output)
    if template.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("Generation requires an XLSX template and XLSX output")
    if template.resolve() == output.resolve():
        raise ValueError("Output must not overwrite the template")
    template_bytes = template.read_bytes()
    digest = hashlib.sha256(template_bytes).hexdigest()
    original_parts = package(BytesIO(template_bytes))
    parts = dict(original_parts)
    paths = sheet_paths(parts)
    if profile["sheet"] not in paths:
        raise ValueError(f"Missing worksheet {profile['sheet']!r}")
    values = projected_values(draft, profile, result, index)
    wb = openpyxl.load_workbook(BytesIO(template_bytes), data_only=False)
    sheet = wb[profile["sheet"]]
    for address, value in computed_values(draft, profile, result).items():
        if sheet[address].data_type != "f":
            values[address] = value
    root = parse(parts[paths[profile["sheet"]]])
    for address, value in values.items():
        if isinstance(sheet[address], openpyxl.cell.cell.MergedCell):
            raise ValueError(f"Profile writes to non-anchor merged cell {address}")
        for validation in sheet.data_validations.dataValidation:
            if address in validation.sqref and validation.type == "list" and validation.formula1 and validation.formula1.startswith('"'):
                allowed = validation.formula1.strip('"').split(",")
                if value not in allowed:
                    raise ValueError(f"{address}: {value!r} not in template validation list {allowed}")
        write_value(cell_node(root, address), value)
        sheet[address] = float(value) if hasattr(value, "as_tuple") else value
    parts[paths[profile["sheet"]]] = serialize(root)
    calculate = evaluator(wb)
    unsupported = []
    formula_count = 0
    caches = {}
    for ws in wb:
        root = parse(parts[paths[ws.title]])
        for row in ws:
            for cell in row:
                if cell.data_type != "f":
                    continue
                formula_count += 1
                try:
                    value = calculate(ws.title, cell.coordinate)
                except (ValueError, KeyError, TypeError, ZeroDivisionError) as error:
                    unsupported.append(f"{ws.title}!{cell.coordinate}: {error}")
                    value = None
                write_cache(cell_node(root, cell.coordinate), value)
                if value is not None:
                    caches[(ws.title, cell.coordinate)] = value
        parts[paths[ws.title]] = serialize(root)
    apply_print(parts, paths, profile)
    added = {"parts": [], "relationship_ids": [], "sheet_names": []}
    if extra_sheets:
        import mapping
        checks = [
            ["Check: template SHA-256", digest],
            ["Check: original package parts", len(original_parts)],
            ["Check: template preservation", "Original formulas, cell styles, drawings/media and relationships checked before publication; only declared values/caches/print settings and registered supplementary sheets may differ."],
            ["Check: original formulas", formula_count],
            ["Check: formula caches", f"{len(caches)} supported results read back before publication; {len(unsupported)} unsupported results."],
            ["Check: unsupported formula caches", "\n".join(unsupported) if unsupported else "None"],
            ["Check: supplementary visibility", "hidden" if hidden_sheets else "visible"],
            ["Check: delivery", "Workbook is self-contained; source locators and rationales are on Allocation ledger. Notes contains metadata, assumptions and band findings."],
        ]
        specs = [(name, list(rows) + checks if name == "Notes" else rows, widths)
                 for name, rows, widths in extra_sheets]
        styles = mapping.template_styles(wb)
        added = mapping.attach(parts, specs, hidden=hidden_sheets, styles=styles)
    output.parent.mkdir(parents=True, exist_ok=True)
    # A staged file is verified before publishing; failures cannot replace prior output.
    import tempfile
    with tempfile.TemporaryDirectory(prefix=".coe-tos-", dir=output.parent) as directory:
        staged = Path(directory) / output.name
        with ZipFile(BytesIO(template_bytes)) as source, ZipFile(staged, "w") as target:
            source_items = source.infolist()
            source_names = {item.filename for item in source_items}
            for item in source_items:
                target.writestr(item, parts[item.filename])
            for name, payload in parts.items():
                if name not in source_names:
                    target.writestr(name, payload)
        fidelity = verify(template, staged, profile, set(values), added, original_parts)
        cached = openpyxl.load_workbook(staged, data_only=True)
        try:
            for address, expected in values.items():
                expected = float(expected) if hasattr(expected, "as_tuple") else expected
                actual = cached[profile["sheet"]][address].value
                matches = actual in (None, "") if expected in (None, "") else actual == expected
                if not matches:
                    raise ValueError(f"Input read-back differs: {address}")
            for (sheet_name, address), expected in caches.items():
                actual = cached[sheet_name][address].value
                if (actual if actual is not None else "") != expected:
                    raise ValueError(f"Formula cache read-back differs: {sheet_name}!{address}")
            for ws in cached:
                if any(cell.data_type == "e" for row in ws for cell in row):
                    raise ValueError(f"Workbook contains Excel error cells: {ws.title}")
        finally:
            cached.close()
        if hashlib.sha256(template.read_bytes()).hexdigest() != digest:
            raise ValueError("Source template changed during build")
        staged.replace(output)
    fidelity.update(template_sha256=digest, formula_count=formula_count,
                    unsupported_formula_caches=unsupported,
                    formula_caches_verified=len(caches),
                    supplementary_visibility="hidden" if hidden_sheets else "visible")
    return fidelity


def preview_workbook(template, output, profile):
    """Prepare a blank template copy with the same explicit printing profile."""
    template, output = Path(template), Path(output)
    if template.resolve() == output.resolve():
        raise ValueError("Preview copy must not overwrite the template")
    parts = package(template)
    apply_print(parts, sheet_paths(parts), profile)
    with ZipFile(template) as source, ZipFile(output, "w") as target:
        for item in source.infolist():
            target.writestr(item, parts[item.filename])
    verify(template, output, profile, set())
