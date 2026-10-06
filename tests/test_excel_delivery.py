"""Excel-only delivery keeps the evidence/checks without extra output files."""
import contextlib
from decimal import Decimal
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "coe-tos" / "scripts"))
from coe_tos import main
from mapping import worksheet_xml
from lxml import etree


class ExcelDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.example = ROOT / "examples" / "draft.json"
        self.draft = json.loads(self.example.read_text())
        self.out = self.root / "out"

    def build(self, *flags, draft=None):
        command = ["build", "--draft", str(draft or self.example), "--out", str(self.out), *flags]
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return main(command)

    def test_default_build_writes_only_one_workbook_without_office(self):
        with patch("export_pdf.export", side_effect=AssertionError("Office must not run")) as export, \
                patch("export_pdf.preflight", side_effect=AssertionError("Office preflight must not run")) as preflight:
            self.assertEqual(self.build(), 0)
            export.assert_not_called()
            preflight.assert_not_called()
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), ["TOS.xlsx"])
        book = openpyxl.load_workbook(self.out / "TOS.xlsx", data_only=True)
        self.assertEqual(book.sheetnames, ["Table of Specifications", "Assessment mapping", "Topic allocation", "Allocation ledger", "Notes"])
        self.assertTrue(all(ws.sheet_state == "visible" for ws in book))
        self.assertEqual([book.worksheets[0][a].value for a in ("D23", "G23", "J23", "L23")], [16, 24, 60, 100])

    def test_sources_and_locators_are_inside_the_workbook(self):
        self.build()
        book = openpyxl.load_workbook(self.out / "TOS.xlsx", data_only=True)
        ledger = book["Allocation ledger"]
        headers = [cell.value for cell in ledger[5]]
        source_col, locator_col = headers.index("Sources"), headers.index("Source locators")
        rows = {row[0].value: [cell.value for cell in row] for row in ledger.iter_rows(min_row=6)}
        for entry in self.draft["allocations"]:
            row = rows[entry["id"]]
            for evidence in entry["evidence"]:
                self.assertIn(evidence["source"], row[source_col])
                self.assertIn(f"{evidence['source']}: {evidence['locator']}", row[locator_col])
            self.assertEqual(row[9], entry["rationale"])
        self.assertTrue(ledger.cell(6, locator_col + 1).alignment.wrap_text)

    def test_notes_embed_metadata_bands_and_preservation_results(self):
        self.build()
        book = openpyxl.load_workbook(self.out / "TOS.xlsx", data_only=True)
        notes = {row[0].value: row[1].value for row in book["Notes"].iter_rows(min_row=6) if row[0].value}
        self.assertEqual(notes["Metadata: course_title"], "Numerical Methods")
        self.assertIn("10-20%", notes["Cognitive: remembering"])
        self.assertEqual(notes["Check: original formulas"], 58)
        self.assertEqual(notes["Check: unsupported formula caches"], "None")
        self.assertIn("58 supported", notes["Check: formula caches"])
        self.assertIn("Not performed", notes["Output visual review"])

    def test_legacy_xlsx_only_flag_is_a_single_file_delivery(self):
        self.build("--xlsx-only", "--with-mapping")
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), ["TOS.xlsx"])

    def test_diagnostics_are_explicit(self):
        self.build("--diagnostics")
        names = {p.name for p in self.out.iterdir()}
        self.assertIn("TOS.xlsx", names)
        self.assertTrue({"review.md", "validation.json", "fidelity.json", "mapping.csv", "timings.json"} <= names)
        self.assertNotIn("TOS.pdf", names)

    def test_explicit_pdf_preserves_form_only_export_visibility(self):
        def export_stub(source, output, *args):
            book = openpyxl.load_workbook(source)
            self.assertTrue(all(ws.sheet_state == "hidden" for ws in book.worksheets[1:]))
            Path(output).write_bytes(b"native backend test stub")
            return {"backend": "test", "pages": 1}
        with patch("export_pdf.preflight", return_value="test"), patch("export_pdf.export", side_effect=export_stub):
            self.assertEqual(self.build("--pdf"), 0)
        self.assertEqual({p.name for p in self.out.iterdir()}, {"TOS.xlsx", "TOS.pdf"})

    def test_failure_never_overwrites_a_previous_workbook(self):
        self.build()
        original = (self.out / "TOS.xlsx").read_bytes()
        invalid = self.root / "invalid.json"
        self.draft["total_points"] = 101
        invalid.write_text(json.dumps(self.draft))
        with self.assertRaises(SystemExit):
            self.build(draft=invalid)
        self.assertEqual((self.out / "TOS.xlsx").read_bytes(), original)
        self.assertEqual({p.name for p in self.out.iterdir()}, {"TOS.xlsx"})

    def test_unsupported_formula_is_disclosed_without_guessing(self):
        self.draft["metadata"]["course_title"] = "Unsupported formula test"
        draft = self.root / "draft.json"
        draft.write_text(json.dumps(self.draft))
        template = self.root / "custom.xlsx"
        book = openpyxl.Workbook()
        book.active.title = "Custom"
        book.active["A1"] = "=UNSUPPORTEDFUNC(1)"
        book.save(template)
        profile = json.loads((ROOT / "skills/coe-tos/assets/cjc-profile.json").read_text())
        profile.update(sheet="Custom", fields={}, prefixes={}, print_area=None, fit_to_page=False)
        profile_path = self.root / "profile.json"
        profile_path.write_text(json.dumps(profile))
        self.assertEqual(self.build("--template", str(template), "--profile", str(profile_path), draft=draft), 0)
        result = openpyxl.load_workbook(self.out / "TOS.xlsx", data_only=True)
        self.assertIsNone(result["Custom"]["A1"].value)
        notes = {row[0].value: row[1].value for row in result["Notes"].iter_rows(min_row=6) if row[0].value}
        self.assertIn("Unsupported formula", notes["Check: unsupported formula caches"])

    def test_decimal_sheet_values_and_wide_columns_are_valid(self):
        from openpyxl.utils import get_column_letter
        xml = etree.fromstring(worksheet_xml([[Decimal("1.234567890123456789"), *range(60)]], []))
        cells = xml.findall(".//{*}c")
        self.assertEqual(cells[-1].get("r"), f"{get_column_letter(61)}1")
        self.assertEqual(cells[0].find("{*}v").text, "1.234567890123456789")


if __name__ == "__main__":
    unittest.main()
