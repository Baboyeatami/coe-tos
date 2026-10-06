import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "coe-tos" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import openpyxl
from lxml import etree

from coe_tos import main
from mapping import sheets as mapping_sheets
from scoring import validate
from workbooks import package, sheet_paths, verify

ASSETS = ROOT / "skills" / "coe-tos" / "assets"
EXAMPLE = ROOT / "examples" / "draft.json"


class SupplementarySheetTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.template = ASSETS / "cjc-template.xlsx"
        self.draft = json.loads(EXAMPLE.read_text())
        self.profile = json.loads((ASSETS / "cjc-profile.json").read_text())
        self.result = validate(self.draft, self.profile)

    def build(self, name, *flags):
        out = self.directory / name
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            status = main(["build", "--draft", str(EXAMPLE), "--out", str(out), "--xlsx-only", *flags])
        self.assertEqual(status, 0)
        return out

    def test_default_build_adds_no_sheet(self):
        out = self.build("plain")
        workbook = openpyxl.load_workbook(out / "TOS.xlsx")
        self.assertEqual(workbook.sheetnames, ["Table of Specifications"])
        self.assertEqual(json.loads((out / "fidelity.json").read_text())["added_sheets"], [])

    def test_with_mapping_adds_hidden_sheets_that_reconcile(self):
        out = self.build("mapped", "--with-mapping")
        workbook = openpyxl.load_workbook(out / "TOS.xlsx", data_only=True)
        self.assertEqual(workbook.sheetnames,
                         ["Table of Specifications", "Assessment mapping", "Topic allocation",
                          "Allocation ledger", "Notes"])
        for name in workbook.sheetnames[1:]:
            self.assertEqual(workbook[name].sheet_state, "hidden", name)

        mapping = workbook["Assessment mapping"]
        headers = [cell.value for cell in mapping[5]]
        self.assertEqual(headers[:4], ["Ref", "Published criterion", "Source points", "Ledger points"])
        rows = {row[0].value: row for row in mapping.iter_rows(min_row=6, max_row=mapping.max_row)}
        criterion = rows["Q1-C"]
        self.assertEqual([criterion[i].value for i in (2, 3, 4, 5, 6)], [10, 10, 2, 4, 4])
        self.assertIn("Published subcriteria", criterion[7].value)
        total = rows["TOTAL"]
        expected = [self.result["total"]] * 2 + [self.result["cognitive"][group]
                                                 for group in ("remembering", "understanding", "thinking")]
        self.assertEqual([total[i].value for i in (2, 3, 4, 5, 6)], [int(value) for value in expected])

        topics = workbook["Topic allocation"]
        topic_rows = [[cell.value for cell in row] for row in topics.iter_rows(min_row=6, max_row=topics.max_row)]
        topic_total = [row for row in topic_rows if row[0] == "TOTAL"][0]
        self.assertEqual([topic_total[i] for i in (1, 2, 3, 4)],
                         [int(self.result["cognitive"][group]) for group in
                          ("remembering", "understanding", "thinking")] + [int(self.result["total"])])
        self.assertTrue(str(topic_total[5]).endswith("%"))

        ledger = workbook["Allocation ledger"]
        ledger_rows = [[cell.value for cell in row] for row in ledger.iter_rows(min_row=6, max_row=ledger.max_row)]
        ledger_total = [row for row in ledger_rows if row[0] == "TOTAL"][0]
        self.assertEqual(ledger_total[5], int(self.result["total"]))
        self.assertEqual(sum(row[5] for row in ledger_rows if row[0] != "TOTAL"), int(self.result["total"]))

    def test_supplementary_sheets_preserve_the_template_package(self):
        out = self.build("fidelity", "--with-mapping")
        before, after = package(self.template), package(out / "TOS.xlsx")
        added = set(json.loads((out / "fidelity.json").read_text())["added_parts"])
        self.assertEqual(set(after) - set(before), added)
        # sheet1.xml, workbook.xml, content types and workbook rels legitimately change;
        # verify() checks those structurally. Everything else must stay byte-identical.
        editable = {"xl/workbook.xml", "[Content_Types].xml", "xl/_rels/workbook.xml.rels"}
        for name, sheet in sheet_paths(before).items():
            if name == self.profile["sheet"]:
                editable.add(sheet)
        for name in before:
            if name in editable:
                continue
            self.assertEqual(before[name], after[name], name)
        self.assertEqual(hash_bytes(self.template), json.loads((out / "fidelity.json").read_text())["template_sha256"])

    def test_undeclared_sheet_addition_is_refused_by_the_fidelity_gate(self):
        out = self.build("gate", "--with-mapping")
        with self.assertRaisesRegex(ValueError, "added or dropped"):
            verify(self.template, out / "TOS.xlsx", self.profile, set())

    def test_visible_flag_shows_the_sheets(self):
        out = self.build("visible", "--with-mapping", "--mapping-visible")
        workbook = openpyxl.load_workbook(out / "TOS.xlsx")
        self.assertEqual(workbook["Assessment mapping"].sheet_state, "visible")

    def test_sheet_names_are_not_overwritten_by_a_template_sheet(self):
        collision = dict(self.draft)
        collision["topics"] = collision["topics"][:2]
        specification = mapping_sheets(collision, self.result)
        parts = package(self.template)
        from mapping import attach
        added = attach(parts, specification, hidden=True)
        self.assertEqual(len(set(added["sheet_names"])), len(added["sheet_names"]))
        self.assertEqual(len(set(added["relationship_ids"])), len(added["relationship_ids"]))


    def test_supplementary_sheets_carry_a_print_setup(self):
        # Without a page setup these sheets print across many pages with split columns.
        out = self.build("printsetup", "--with-mapping")
        parts = package(out / "TOS.xlsx")
        added = json.loads((out / "fidelity.json").read_text())["added_parts"]
        self.assertEqual(len(added), 4)
        for part in added:
            root = etree.fromstring(parts[part])
            setup = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}pageSetup")
            self.assertIsNotNone(setup, part)
            self.assertEqual(setup.get("orientation"), "landscape", part)
            self.assertEqual(setup.get("fitToWidth"), "1", part)
            self.assertEqual(setup.get("fitToHeight"), "0", part)
            prop = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetPr")
            self.assertIsNotNone(prop, part)
            self.assertIsNotNone(prop.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}pageSetUpPr"), part)

    def test_ref_column_is_wide_enough_for_subcriterion_labels(self):
        out = self.build("widths", "--with-mapping")
        parts = package(out / "TOS.xlsx")
        added = json.loads((out / "fidelity.json").read_text())["added_parts"]
        part = next(p for p in added if "Assessment mapping" in workbook_sheet_name(parts, p))
        root = etree.fromstring(parts[part])
        first = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}cols")[0]
        self.assertGreaterEqual(float(first.get("width")), 24)


def workbook_sheet_name(parts, part):
    workbook = etree.fromstring(parts["xl/workbook.xml"])
    rels = {r.get("Id"): r.get("Target") for r in etree.fromstring(parts["xl/_rels/workbook.xml.rels"])}
    for sheet in workbook.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets"):
        if rels[sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")].endswith(part.split("/")[-1]):
            return sheet.get("name")
    return ""


def hash_bytes(path):
    import hashlib
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


if __name__ == "__main__":
    unittest.main()