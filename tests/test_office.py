"""Platform discovery, backend selection and legacy-input conversion contracts."""
import contextlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/coe-tos/scripts"))

import documents
import export_pdf
import office
from coe_tos import main
from docx import Document
import openpyxl
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def write_pdf(path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({
        NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 40 750 Td (Verified TOS export) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with Path(path).open("wb") as target:
        writer.write(target)


class DiscoveryTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)

    def binary(self, name):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"discovery fixture, never executed")
        return path

    def test_override_wins_and_invalid_override_is_not_silently_ignored(self):
        override = self.binary("custom folder/soffice.exe")
        with patch.dict(os.environ, {"COE_TOS_SOFFICE": str(override)}, clear=True), \
                patch.object(office.shutil, "which") as which:
            self.assertEqual(office.office_binary(), str(override.resolve()))
            which.assert_not_called()
            os.environ["COE_TOS_SOFFICE"] = str(self.root / "missing.exe")
            with self.assertRaisesRegex(ValueError, "COE_TOS_SOFFICE"):
                office.office_binary()

    def test_windows_default_locations_work_without_path(self):
        for variable in ("ProgramW6432", "ProgramFiles", "ProgramFiles(x86)"):
            with self.subTest(variable=variable):
                root = self.root / variable
                binary = self.binary(f"{variable}/LibreOffice/program/soffice.exe")
                with patch.dict(os.environ, {variable: str(root)}, clear=True), \
                        patch.object(office.platform, "system", return_value="Windows"), \
                        patch.object(office.shutil, "which", return_value=None):
                    self.assertEqual(office.office_binary(), str(binary))

    def test_path_takes_precedence_over_default_install_location(self):
        binary = self.binary("on path/soffice.exe")
        self.binary("Program Files/LibreOffice/program/soffice.exe")
        with patch.dict(os.environ, {"ProgramFiles": str(self.root / "Program Files")}, clear=True), \
                patch.object(office.platform, "system", return_value="Windows"), \
                patch.object(office.shutil, "which", return_value=str(binary)):
            self.assertEqual(office.office_binary(), str(binary))

    def test_non_windows_checks_do_not_import_com(self):
        with patch.object(office.platform, "system", return_value="Darwin"), \
                patch.object(office.importlib, "import_module") as imports:
            self.assertIn("requires Windows", office.windows_office_error())
            imports.assert_not_called()

    def test_missing_pywin32_gives_install_command_without_starting_office(self):
        with patch.object(office.platform, "system", return_value="Windows"), \
                patch.object(office.importlib, "import_module", side_effect=ImportError("pythoncom")):
            self.assertIn("requirements-windows-office.txt", office.windows_office_error())

    def test_registration_in_other_bitness_view_is_found_read_only(self):
        registry = Mock(KEY_READ=1, KEY_WOW64_64KEY=2, KEY_WOW64_32KEY=4)
        key = Mock()
        key.__enter__ = Mock(return_value=key)
        key.__exit__ = Mock(return_value=False)
        registry.OpenKey.side_effect = [FileNotFoundError(), FileNotFoundError(), key]
        registry.QueryValueEx.return_value = ("{Excel-CLSID}", 1)
        with patch.object(office.platform, "system", return_value="Windows"), \
                patch.object(office.importlib, "import_module", side_effect=[Mock(), Mock(), registry]):
            self.assertIsNone(office.windows_office_error())
        self.assertEqual(registry.OpenKey.call_args.args[1], "Excel.Application\\CLSID")
        self.assertEqual(registry.OpenKey.call_args.args[3], 5)


class BackendTests(unittest.TestCase):
    def test_auto_prefers_windows_excel_and_explicit_backend_never_falls_back(self):
        with patch.object(export_pdf, "windows_office_error", return_value=None), \
                patch.object(documents, "office_binary") as libreoffice:
            self.assertEqual(export_pdf.preflight(), "excel-windows")
            libreoffice.assert_not_called()
        with patch.object(export_pdf, "windows_office_error", return_value="pywin32 missing"), \
                patch.object(documents, "office_binary") as libreoffice:
            with self.assertRaisesRegex(ValueError, "pywin32 missing"):
                export_pdf.preflight("excel-windows")
            libreoffice.assert_not_called()

    def test_auto_can_use_libreoffice_when_windows_excel_is_unavailable(self):
        with patch.object(export_pdf.platform, "system", return_value="Windows"), \
                patch.object(export_pdf, "windows_office_error", return_value="Excel not registered"), \
                patch.object(documents, "office_binary", return_value="soffice.exe"):
            self.assertEqual(export_pdf.preflight(), "libreoffice")

    def test_existing_mac_excel_selection_is_retained(self):
        with patch.object(export_pdf.platform, "system", return_value="Darwin"), \
                patch.object(export_pdf.Path, "exists", return_value=True), \
                patch.object(export_pdf, "windows_office_error", return_value="requires Windows"):
            self.assertEqual(export_pdf.preflight(), "excel-mac")

    def test_missing_windows_backends_reports_how_to_use_excel_only(self):
        with patch.object(export_pdf.platform, "system", return_value="Windows"), \
                patch.object(export_pdf, "windows_office_error", return_value="pywin32 missing"), \
                patch.object(documents, "office_binary", return_value=None):
            with self.assertRaisesRegex(ValueError, "Omit --pdf.*pywin32 missing"):
                export_pdf.preflight()

    def test_excel_export_uses_copy_and_failed_export_preserves_previous_pdf(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, output = root / "original.xlsx", root / "previous.pdf"
            source.write_bytes(b"verified workbook")
            output.write_bytes(b"previous PDF")

            def exporter(operation, copied, generated):
                self.assertNotEqual(copied, source)
                self.assertEqual(copied.read_bytes(), b"verified workbook")
                write_pdf(generated)

            with patch.object(export_pdf, "preflight", return_value="excel-windows"), \
                    patch("windows_office.run_office", side_effect=exporter), \
                    patch.object(documents, "convert_office") as fallback:
                info = export_pdf.export(source, output)
                self.assertEqual(info["backend"], "excel-windows")
                self.assertEqual(info["pages"], 1)
                self.assertTrue(info["searchable_text"])
                previous = output.read_bytes()
                with patch("windows_office.run_office", side_effect=ValueError("Excel export failed")):
                    with self.assertRaisesRegex(ValueError, "Excel export failed"):
                        export_pdf.export(source, output)
                fallback.assert_not_called()
            self.assertEqual(source.read_bytes(), b"verified workbook")
            self.assertEqual(output.read_bytes(), previous)

    def test_windows_backend_is_accepted_by_every_pdf_cli_command(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)

            def exporter(source, output, *args):
                write_pdf(output)
                return {"backend": "excel-windows", "pages": 1}

            commands = [
                ["build", "--draft", str(ROOT / "examples/draft.json"), "--out", str(root / "build"), "--pdf"],
                ["preview-template", "--out", str(root / "blank.pdf")],
                ["export-pdf", str(ROOT / "skills/coe-tos/assets/cjc-template.xlsx"), "--out", str(root / "form.pdf")],
            ]
            with patch.object(export_pdf, "preflight", return_value="excel-windows"), \
                    patch.object(export_pdf, "export", side_effect=exporter) as export, \
                    contextlib.redirect_stdout(io.StringIO()):
                for command in commands:
                    with self.subTest(command=command[0]):
                        self.assertEqual(main([*command, "--backend", "excel-windows"]), 0)
                        self.assertEqual(export.call_args.args[2], "excel-windows")

    def test_invalid_native_pdf_preserves_a_previous_valid_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, output = root / "source.xlsx", root / "previous.pdf"
            source.write_bytes(b"verified workbook")
            write_pdf(output)
            previous = output.read_bytes()

            def empty_export(operation, copied, generated):
                with generated.open("wb") as stream:
                    PdfWriter().write(stream)

            with patch.object(export_pdf, "preflight", return_value="excel-windows"), \
                    patch("windows_office.run_office", side_effect=empty_export):
                with self.assertRaisesRegex(ValueError, "empty"):
                    export_pdf.export(source, output)
            self.assertEqual(output.read_bytes(), previous)


class LegacyConversionTests(unittest.TestCase):
    def test_native_converter_matches_source_type_and_keeps_original_name(self):
        for suffix, application in ((".doc", "Word"), (".xls", "Excel")):
            with self.subTest(suffix=suffix), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / ("épreuve with spaces" + suffix)
                source.write_bytes(b"legacy source")

                def converter(operation, copied, target):
                    self.assertNotEqual(copied, source)
                    self.assertEqual(copied.read_bytes(), b"legacy source")
                    if operation == "docx":
                        doc = Document()
                        doc.add_paragraph("Q1: 50 points")
                        doc.save(target)
                    else:
                        book = openpyxl.Workbook()
                        book.active["A1"] = "Q1: 50 points"
                        book.save(target)
                    return target

                with patch.object(office, "windows_office_error", return_value=None), \
                        patch("windows_office.run_office", side_effect=converter), \
                        patch.object(documents, "convert_office") as fallback:
                    result = documents.extract(source)
                    fallback.assert_not_called()
                self.assertEqual(result["source"], source.name)
                self.assertTrue(any("Q1: 50" in entry["text"] for entry in result["segments"]))
                self.assertTrue(any(f"Microsoft {application}" in text for text in result["warnings"]))
                self.assertEqual(source.read_bytes(), b"legacy source")

    def test_conversion_failure_is_not_hidden_by_another_engine(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "rubric.doc"
            source.write_bytes(b"source")
            with patch.object(office, "windows_office_error", return_value=None), \
                    patch("windows_office.run_office", side_effect=ValueError("Word conversion failed")), \
                    patch.object(documents, "convert_office") as fallback:
                with self.assertRaisesRegex(ValueError, "Word conversion failed"):
                    documents.extract(source)
                fallback.assert_not_called()

    def test_non_office_formats_keep_libreoffice_and_can_be_explicitly_selected(self):
        with patch.object(office, "windows_office_error", return_value=None), \
                patch.object(office, "office_binary", return_value="soffice"):
            for suffix in (".odt", ".ods", ".rtf"):
                self.assertEqual(office.input_backend(suffix), "libreoffice")
                with self.assertRaisesRegex(ValueError, "supports .doc and .xls"):
                    office.input_backend(suffix, "ms-office")
            self.assertEqual(office.input_backend(".doc", "libreoffice"), "libreoffice")


if __name__ == "__main__":
    unittest.main()
