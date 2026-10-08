import copy
import json
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "coe-tos" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import openpyxl
from docx import Document
from lxml import etree as ET
from pypdf import PdfWriter
from documents import extract
from scoring import validate
from workbooks import build, package, projected_values, verify, preview_workbook
from formulas import evaluator
from coe_tos import main

ASSETS = ROOT / "skills" / "coe-tos" / "assets"


class RefFittingTests(unittest.TestCase):
    def test_long_reference_lists_abbreviate_instead_of_clipping(self):
        from workbooks import fit_refs
        self.assertEqual(fit_refs(["Q1-c1b", "Q2-c2b", "Q3-c3a"], 16), "Q1-c1b +2")
        self.assertEqual(fit_refs(["Q1-a", "Q2-b", "Q3-c", "Q4-d"], 16), "Q1-a; Q2-b +2")

    def test_short_reference_lists_are_untouched(self):
        from workbooks import fit_refs
        self.assertEqual(fit_refs(["EX-c1", "EX-c2"], 16), "EX-c1; EX-c2")
        self.assertEqual(fit_refs(["Q3-c3b"], 16), "Q3-c3b")
        self.assertEqual(fit_refs([], 16), "")
        self.assertEqual(fit_refs(["Q1-a", "Q2-b", "Q3-c"], None), "Q1-a; Q2-b; Q3-c")

    def test_project_never_exceeds_the_declared_reference_width(self):
        draft = json.loads((ROOT / "examples" / "draft.json").read_text())
        profile = json.loads((ASSETS / "cjc-profile.json").read_text())
        result = validate(draft, profile)
        limit = profile["refs_max_chars"]
        values = projected_values(draft, profile, result)
        column = profile["columns"]["understanding_refs"]
        for index in range(len(draft["topics"])):
            cell = values.get(f"{column}{profile['first_row'] + index}")
            if cell:
                self.assertLessEqual(len(cell), limit, cell)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.draft = json.loads((ROOT / "examples" / "draft.json").read_text())
        self.profile = json.loads((ASSETS / "cjc-profile.json").read_text())
        self.temporary = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def test_example_reconciles_all_levels(self):
        result = validate(self.draft, self.profile)
        self.assertEqual(result["errors"], [])
        self.assertEqual(list(result["cognitive"].values()), [16, 24, 60])
        self.assertEqual(list(result["question_scores"]["Q1"].values()), [8, 12, 30])
        self.assertTrue(all(result["bands"].values()))

    def test_preserves_subcriteria_even_when_overall_total_matches(self):
        # Same criterion and exam totals, but the original 4/4 split changed to 5/3.
        self.draft["allocations"][8]["points"] = 5
        self.draft["allocations"][8]["scores"]["understanding"] = 5
        self.draft["allocations"][9]["points"] = 3
        self.draft["allocations"][9]["scores"]["thinking"] = 3
        errors = validate(self.draft, self.profile)["errors"]
        self.assertTrue(any("published subcriterion" in error for error in errors))
        self.assertFalse(any("total_points" in error for error in errors))

    def test_duplicate_allocation_rejected(self):
        self.draft["allocations"].append(copy.deepcopy(self.draft["allocations"][0]))
        self.assertTrue(validate(self.draft, self.profile)["errors"])

    def test_missing_evidence_rejected(self):
        self.draft["allocations"][0]["evidence"] = []
        self.assertTrue(any("evidence" in e for e in validate(self.draft, self.profile)["errors"]))

    def test_nonfinite_and_negative_scores_rejected(self):
        for value in ("NaN", "Infinity", -1, True):
            with self.subTest(value=value):
                self.draft["allocations"][0]["scores"]["remembering"] = value
                self.assertTrue(validate(self.draft, self.profile)["errors"])

    def test_capacity_rejected_without_silent_truncation(self):
        self.draft["topics"].append({"id": "extra", "title": "Extra"})
        self.assertTrue(any("capacity" in e for e in validate(self.draft, self.profile)["errors"]))

    def test_band_conflict_is_explicit_not_forced(self):
        for allocation in self.draft["allocations"]:
            allocation["scores"]["understanding"] += allocation["scores"]["remembering"]
            allocation["scores"]["remembering"] = 0
        result = validate(self.draft, self.profile)
        self.assertFalse(result["errors"])
        self.assertFalse(result["bands"]["remembering"])
        self.assertTrue(any("outside" in w for w in result["warnings"]))

    def test_formula_caches_and_binary_template_fidelity(self):
        output = self.directory / "TOS.xlsx"
        template = ASSETS / "cjc-template.xlsx"
        original = template.read_bytes()
        result = validate(self.draft, self.profile)
        fidelity = build(template, output, self.draft, self.profile, result)
        self.assertEqual(template.read_bytes(), original)
        self.assertEqual(fidelity["unsupported_formula_caches"], [])
        before, after = package(template), package(output)
        for part in before:
            if part not in ("xl/workbook.xml", "xl/worksheets/sheet1.xml"):
                self.assertEqual(before[part], after[part], part)
        values = openpyxl.load_workbook(output, data_only=True)
        formulas = openpyxl.load_workbook(output, data_only=False)
        main_sheet = values[self.profile["sheet"]]
        self.assertEqual([main_sheet[a].value for a in ("D23", "G23", "J23", "L23", "M23")], [16, 24, 60, 100, 1])
        calculate = evaluator(formulas)
        for ws in formulas:
            for row in ws:
                for cell in row:
                    if cell.data_type == "f":
                        actual = values[ws.title][cell.coordinate].value
                        expected = calculate(ws.title, cell.coordinate)
                        if expected == "":
                            self.assertIn(actual, (None, ""))
                        else:
                            self.assertAlmostEqual(actual, expected, places=10)

    def test_formula_overwrite_and_invalid_validation_refused(self):
        result = validate(self.draft, self.profile)
        output = self.directory / "bad.xlsx"
        self.profile["fields"]["course_title"] = "L23"
        with self.assertRaisesRegex(ValueError, "formula"):
            build(ASSETS / "cjc-template.xlsx", output, self.draft, self.profile, result)
        self.assertFalse(output.exists())
        self.profile["fields"]["course_title"] = "A9"
        self.draft["metadata"]["assessment"] = "Unknown exam"
        with self.assertRaisesRegex(ValueError, "validation"):
            build(ASSETS / "cjc-template.xlsx", output, self.draft, self.profile, result)

    def test_style_tampering_detected(self):
        output = self.directory / "TOS.xlsx"
        template = ASSETS / "cjc-template.xlsx"
        result = validate(self.draft, self.profile)
        build(template, output, self.draft, self.profile, result)
        parts = package(output)
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        root.find(".//{*}c[@r='B13']").set("s", "0")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
        with ZipFile(output, "w") as target:
            for name, data in parts.items():
                target.writestr(name, data)
        with self.assertRaisesRegex(ValueError, "formatting"):
            verify(template, output, self.profile, set(projected_values(self.draft, self.profile, result)))

    def test_different_template_profile(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Custom Form"
        template = self.directory / "custom.xlsx"
        wb.save(template)
        self.profile.update(sheet=ws.title, fields={}, prefixes={}, print_area=None, fit_to_page=False)
        output = self.directory / "custom-output.xlsx"
        result = validate(self.draft, self.profile)
        build(template, output, self.draft, self.profile, result)
        actual = openpyxl.load_workbook(output).active
        self.assertEqual(actual["B13"].value, "Governing model and definitions")
        self.assertEqual(actual["L23"].value, 100)

    def test_blank_preview_only_changes_printing_and_calculation(self):
        output = self.directory / "blank.xlsx"
        template = ASSETS / "cjc-template.xlsx"
        original = template.read_bytes()
        preview_workbook(template, output, self.profile)
        self.assertEqual(template.read_bytes(), original)
        workbook = openpyxl.load_workbook(output)
        self.assertIn("A1", str(workbook.active.print_area).replace("$", ""))
        self.assertEqual(workbook.active.page_setup.fitToHeight, 1)
        self.assertTrue(verify(template, output, self.profile, set())["formula_and_style_fidelity"])

    def test_extraction_word_text_excel_and_scanned_pdf(self):
        path = self.directory / "rubric.docx"
        doc = Document()
        doc.add_heading("Rubric", 1)
        doc.add_table(rows=1, cols=2).cell(0, 0).text = "Q1 = 50 points"
        doc.save(path)
        result = extract(path)
        self.assertTrue(any("Q1 = 50" in s["text"] and "table" in s["locator"] for s in result["segments"]))
        text = self.directory / "rubric.txt"
        text.write_text("Q1: 50\nQ2: 50", encoding="utf-16")
        self.assertEqual(extract(text)["segments"][1]["locator"], "line 2")
        pdf = self.directory / "scan.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        with pdf.open("wb") as stream:
            writer.write(stream)
        self.assertTrue(any("OCR" in w or "--ocr" in w for w in extract(pdf)["warnings"]))
        self.assertTrue(extract(ASSETS / "cjc-template.xlsx")["segments"])

    def test_partial_extraction_has_nonzero_status(self):
        text = self.directory / "ok.txt"
        text.write_text("Readable source")
        binary = self.directory / "bad.bin"
        binary.write_bytes(b"\x00\xfe\x81")
        out = self.directory / "sources.json"
        self.assertEqual(main(["extract", str(text), str(binary), "--out", str(out)]), 1)
        docs = json.loads(out.read_text())["documents"]
        self.assertTrue(docs[0]["segments"])
        self.assertIn("error", docs[1])

    def test_packaged_and_installed_skill_is_self_contained(self):
        archive = self.directory / "coe-tos.zip"
        subprocess.run([sys.executable, str(ROOT / "scripts/package_skill.py"), "--out", str(archive)], check=True, capture_output=True)
        with ZipFile(archive) as package_zip:
            self.assertIsNone(package_zip.testzip())
            for name in ("SKILL.md", "LICENSE", "assets/cjc-template.xlsx", "scripts/coe_tos.py", "scripts/requirements.txt",
                         "scripts/office.py", "scripts/windows_office.py", "scripts/requirements-windows-office.txt"):
                self.assertIn("coe-tos/" + name, package_zip.namelist())
        target = self.directory / "installed/coe-tos"
        command = [sys.executable, str(ROOT / "scripts/install.py"), "--harness", "generic", "--destination", str(target)]
        subprocess.run(command, check=True, capture_output=True)
        self.assertTrue((target / "LICENSE").exists())
        process = subprocess.run([sys.executable, str(target / "scripts/coe_tos.py"), "build", "--draft", str(ROOT / "examples/draft.json"),
                                  "--out", str(self.directory / "standalone-output"), "--xlsx-only"], cwd=self.directory, capture_output=True, text=True)
        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertNotEqual(subprocess.run(command, capture_output=True).returncode, 0)


if __name__ == "__main__":
    unittest.main()
